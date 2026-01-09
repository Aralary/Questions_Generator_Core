from pathlib import Path
from typing import List, Optional
import logging

logger = logging.getLogger(__name__)


def extract_text_from_pdf(pdf_path: Path) -> str:
    try:
        import pdfplumber
        
        text = ""
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        
        logger.info(f"  PDF: {pdf_path.name} — {len(text)} символов")
        return text.strip()
    
    except ImportError:
        logger.error("Установите pdfplumber: pip install pdfplumber")
        return ""
    except Exception as e:
        logger.error(f"Ошибка при чтении PDF {pdf_path.name}: {e}")
        return ""


def extract_text_from_pptx(pptx_path: Path, include_notes: bool = True) -> str:
    try:
        from pptx import Presentation
        
        prs = Presentation(pptx_path)
        all_text = []
        
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = f"--- Слайд {slide_num} ---\n"
            
            # Извлекаем текст из shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text += shape.text + "\n"
                
                # Извлекаем текст из таблиц
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        slide_text += row_text + "\n"
            
            # Заметки докладчика
            if include_notes and slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text
                if notes_text.strip():
                    slide_text += f"\n[Заметки]: {notes_text}\n"
            
            all_text.append(slide_text)
        
        result = "\n\n".join(all_text)
        logger.info(f"  PPTX: {pptx_path.name} — {len(result)} символов ({len(prs.slides)} слайдов)")
        return result.strip()
    
    except ImportError:
        logger.error("Установите python-pptx: pip install python-pptx")
        return ""
    except Exception as e:
        logger.error(f"Ошибка при чтении PPTX {pptx_path.name}: {e}")
        return ""


def load_lectures_for_domain(domain_dir: Path) -> Optional[str]:
    """
    Загружает все лекции из domain/lectures/.
    
    Args:
        domain_dir: Путь к папке домена (например, data/train_data/cryptography/)
        
    Returns:
        Объединённый текст всех лекций или None
    """
    lectures_dir = domain_dir / "lectures"
    
    if not lectures_dir.exists():
        logger.debug(f"Папка с лекциями не найдена: {lectures_dir}")
        return None
    
    all_lectures = []
    
    # Обрабатываем PDF файлы
    pdf_files = list(lectures_dir.glob("*.pdf"))
    for pdf_file in sorted(pdf_files):
        text = extract_text_from_pdf(pdf_file)
        if text:
            all_lectures.append(f"=== {pdf_file.stem} ===\n{text}")
    
    # Обрабатываем PPTX файлы
    pptx_files = list(lectures_dir.glob("*.pptx"))
    for pptx_file in sorted(pptx_files):
        text = extract_text_from_pptx(pptx_file, include_notes=True)
        if text:
            all_lectures.append(f"=== {pptx_file.stem} ===\n{text}")
    
    total_files = len(pdf_files) + len(pptx_files)
    
    if not all_lectures:
        logger.debug(f"Файлы лекций не найдены в {lectures_dir}")
        return None
    
    logger.info(f"  ✓ Загружено {len(all_lectures)} файлов лекций (PDF: {len(pdf_files)}, PPTX: {len(pptx_files)})")
    return "\n\n".join(all_lectures)


def split_text_into_chunks(text: str, max_length: int = 400, overlap: int = 50) -> List[str]:
    """Разбивает текст на смысловые фрагменты с перекрытием."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    
    chunks = []
    current_chunk = ""
    
    for para in paragraphs:
        if len(para) > max_length:
            sentences = para.replace(". ", ".\n").split("\n")
            for sentence in sentences:
                sentence = sentence.strip()
                if not sentence:
                    continue
                    
                if len(current_chunk) + len(sentence) < max_length:
                    current_chunk += sentence + " "
                else:
                    if current_chunk:
                        chunks.append(current_chunk.strip())
                        words = current_chunk.split()
                        current_chunk = " ".join(words[-overlap//10:]) + " " + sentence + " "
                    else:
                        current_chunk = sentence + " "
        else:
            if len(current_chunk) + len(para) < max_length:
                current_chunk += para + "\n\n"
            else:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                    words = current_chunk.split()
                    current_chunk = " ".join(words[-overlap//10:]) + "\n\n" + para + "\n\n"
                else:
                    current_chunk = para + "\n\n"
    
    if current_chunk:
        chunks.append(current_chunk.strip())
    
    return chunks


def filter_relevant_chunks(chunks: List[str], keywords: List[str], top_k: int = 3) -> List[str]:
    """Фильтрует наиболее релевантные фрагменты по ключевым словам."""
    if not chunks or not keywords:
        return chunks[:top_k] if chunks else []
    
    scored_chunks = []
    
    for chunk in chunks:
        score = sum(keyword.lower() in chunk.lower() for keyword in keywords)
        if score > 0:
            scored_chunks.append((score, chunk))
    
    if not scored_chunks:
        return chunks[:top_k]
    
    scored_chunks.sort(reverse=True, key=lambda x: x[0])
    return [chunk for _, chunk in scored_chunks[:top_k]]


def extract_keywords(text: str, top_n: int = 5) -> List[str]:
    """Простое извлечение ключевых слов из текста."""
    stop_words = {
        "и", "в", "на", "с", "для", "как", "что", "это", "или", "по", 
        "из", "к", "о", "от", "у", "за", "при", "до", "об"
    }
    words = text.lower().split()
    
    filtered_words = [
        w.strip(".,?!:;()[]") 
        for w in words 
        if len(w) > 3 and w not in stop_words
    ]
    
    from collections import Counter
    word_freq = Counter(filtered_words)
    
    return [word for word, _ in word_freq.most_common(top_n)]
